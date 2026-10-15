#!/usr/bin/env python3
"""Compara uma planilha (abas HORIZON, NOW...) com os dados/*.json. Não altera nada.

Uso:
    python ferramentas/conferir_planilha.py "caminho\\PLANILHA.xlsx"

Precisa do pacote openpyxl (pip install openpyxl).

Formato esperado de cada aba:
    linha 1: EMPREENDIMENTO: <nome>
    linha 2: R.I <data>
    cabeçalho com: Solicitante | Atividade | Responsável | [Protocolo] | Prazo Ideal | Prazo Real | Motivo do Desvio
    uma linha por atividade, na ordem do processo.

Cada dados/<projeto>.json declara em "_meta"."aba" o nome da aba correspondente.
As atividades do JSON (sem as de referência e sem marcos) devem estar na mesma ordem da planilha.
"""
import glob
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("Instale o openpyxl:  pip install openpyxl")

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def inteiro(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def texto(v):
    return "" if v is None else str(v).strip()


def linhas_da_aba(ws):
    linhas = list(ws.iter_rows(values_only=True))
    topo = next(i for i, r in enumerate(linhas) if texto(r[0]) == "Solicitante")
    cab = [texto(c) for c in linhas[topo]]
    saida = []
    for r in linhas[topo + 1:]:
        if any(c is not None for c in r):
            saida.append({cab[i]: r[i] for i in range(len(cab))})
    return saida


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    problemas = 0
    for caminho in sorted(glob.glob(os.path.join(RAIZ, "dados", "*.json"))):
        dados = json.load(open(caminho, encoding="utf-8"))
        aba = dados.get("_meta", {}).get("aba")
        nome = dados["projeto"]
        if not aba or aba not in wb.sheetnames:
            print(f"[{nome}] aba '{aba}' não encontrada na planilha. Abas: {wb.sheetnames}")
            problemas += 1
            continue
        plan = linhas_da_aba(wb[aba])
        reais = [t for t in dados["tarefas"] if not t.get("referencia") and not t.get("marco")]
        print(f"\n[{nome}]  planilha: {len(plan)} linhas  |  dados: {len(reais)} tarefas")
        if len(plan) != len(reais):
            print("  ! quantidade diferente: confira se entrou ou saiu alguma atividade")
            problemas += 1
        for i, (t, d) in enumerate(zip(reais, plan), 1):
            ideal, real = inteiro(d.get("Prazo Ideal")), inteiro(d.get("Prazo Real"))
            esperado = real if real is not None else ideal
            achados = []
            if esperado != t.get("duracao"):
                achados.append(f"duração: dados={t.get('duracao')} planilha={esperado}")
            if texto(d.get("Motivo do Desvio")) != texto(t.get("motivo")):
                achados.append("motivo do desvio diferente")
            if texto(d.get("Solicitante")) != texto(t.get("solicitante")):
                achados.append(f"solicitante: dados='{t.get('solicitante')}' planilha='{texto(d.get('Solicitante'))}'")
            if inteiro(d.get("Prazo Ideal")) != t.get("prazoIdealPlan"):
                achados.append("prazo ideal da planilha diferente")
            if inteiro(d.get("Prazo Real")) != t.get("prazoRealPlan"):
                achados.append("prazo real da planilha diferente")
            marca = "ok " if not achados else "DIF"
            print(f"  {marca} {i:>2}. {texto(d.get('Atividade'))[:60]}")
            for a in achados:
                print(f"        - {a}")
            problemas += bool(achados)
    print("\nResultado:", "tudo igual" if not problemas else f"{problemas} ponto(s) para revisar")


if __name__ == "__main__":
    main()
