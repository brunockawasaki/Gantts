# -*- coding: utf-8 -*-
"""
Gera o Word "checklist de trabalho" no formato padrão do Bruno: título centralizado + subtítulo,
um cabeçalho por documento-fonte (negrito, azul-marinho, borda inferior), e um item por ponto
[C]/[D]/[C/D] com caixa de seleção "☐" seguida de uma linha em branco com borda fina para
escrever a resposta.

Uso via linha de comando:
    python3 gen_checklist.py --output "Checklist de Correcoes.docx" \
        --title "Checklist de Correções — Regimentos PJC" \
        --subtitle "Pendências [C]/[D] identificadas na revisão de 2026-08-14" \
        --data data.json

data.json:
{
  "documents": [
    {
      "heading": "RE.IG.PJC.001 — Regimento dos Comitês de Due Diligence",
      "items": [
        "Quórum e participantes obrigatórios não definidos na seção 4.",
        "Ata: prazo de 2 dias úteis definido, sem local de armazenamento."
      ]
    }
  ]
}

Ou importe build_checklist() diretamente de outro script Python.
"""
import argparse
import json

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_pbdr_bottom(paragraph, sz, color, space=4):
    """Add a bottom border to a paragraph. Must run before spacing/ind are set on the
    same pPr in some Word versions' validation — inserting as the first child of pPr
    satisfies the CT_PPr schema order here since this template uses neither pStyle nor numPr."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), str(space))
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.insert(0, pBdr)


def _set_arial_10(document):
    style = document.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(10)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:ascii'), 'Arial')
    rFonts.set(qn('w:hAnsi'), 'Arial')
    rFonts.set(qn('w:cs'), 'Arial')


def build_checklist(output_path, title, subtitle, documents):
    """documents: list of {"heading": str, "items": list[str]}"""
    doc = Document()
    _set_arial_10(doc)

    # python-docx's default blank template ships <w:zoom w:val="bestFit"/> without the
    # schema-required w:percent attribute; add it so strict OOXML validation passes.
    _zoom = doc.settings.element.find(qn('w:zoom'))
    if _zoom is not None and _zoom.get(qn('w:percent')) is None:
        _zoom.set(qn('w:percent'), '100')

    sec = doc.sections[0]
    sec.page_width = Pt(595.3)  # A4
    sec.page_height = Pt(841.9)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)

    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(subtitle)
        r.italic = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    for entry in documents:
        heading = doc.add_paragraph()
        heading.paragraph_format.space_before = Pt(18)
        heading.paragraph_format.space_after = Pt(6)
        _set_pbdr_bottom(heading, sz=8, color="1F3864", space=4)
        r = heading.add_run(entry["heading"])
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

        for item_text in entry["items"]:
            item_p = doc.add_paragraph()
            item_p.paragraph_format.space_before = Pt(8)
            item_p.paragraph_format.space_after = Pt(2)
            r = item_p.add_run("☐  " + item_text)
            r.font.size = Pt(10)

            blank_p = doc.add_paragraph()
            blank_p.paragraph_format.space_before = Pt(0)
            blank_p.paragraph_format.space_after = Pt(10)
            _set_pbdr_bottom(blank_p, sz=4, color="AAAAAA", space=4)
            r2 = blank_p.add_run(" ")
            r2.font.size = Pt(10)

    doc.save(output_path)
    return output_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--output", required=True)
    ap.add_argument("--title", required=True)
    ap.add_argument("--subtitle", default="")
    group = ap.add_mutually_exclusive_group(required=True)
    group.add_argument("--data", help="path to a JSON file matching the schema in this file's docstring")
    group.add_argument("--data-json", help="the JSON itself, inline, instead of a file path")
    args = ap.parse_args()

    if args.data:
        with open(args.data, encoding="utf-8") as f:
            payload = json.load(f)
    else:
        payload = json.loads(args.data_json)

    out = build_checklist(args.output, args.title, args.subtitle, payload["documents"])
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
