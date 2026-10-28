# -*- coding: utf-8 -*-
"""Gera o Formulário de Levantamento de Processo (base para POP) — Grupo Gadens.

Uso:
    python gera_formulario_word.py --saida "C:\\caminho\\Formulario.docx" [--logo caminho.png]

Requer: python-docx (pip install python-docx). Validado em jul/2026.
"""
import argparse
import os
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO_PADRAO = os.path.join(HERE, '..', 'assets', 'logo_gadens.png')

# Cores da marca Grupo Gadens (amostradas do logo oficial)
RED, DARK = 'C0272D', '58595B'
MIDGRAY = '767171'      # dicas
TEXT = '333333'         # texto padrão
FILL_LIGHT = 'F2F2F2'   # sombreado claro
FILL_EX = 'F7F7F7'      # linha de exemplo
BORDER = 'B7B7B7'       # bordas de tabela
TAG = 'E3C6C8'          # etiqueta clara sobre barra escura
FULL_DXA = 9638         # 17 cm

def hexrgb(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

# ---------- helpers ----------

def cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def para_shading(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)

def para_left_border(p, color, sz=24, space=4):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), str(sz))
    left.set(qn('w:space'), str(space))
    left.set(qn('w:color'), color)
    pBdr.append(left)
    pPr.append(pBdr)

def para_border(p, side, color, sz=12, space=1):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    b = OxmlElement(f'w:{side}')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), str(sz))
    b.set(qn('w:space'), str(space))
    b.set(qn('w:color'), color)
    pBdr.append(b)
    pPr.append(pBdr)

def tbl_setup(table, col_cms, border_color=BORDER, header_repeat=False):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(FULL_DXA))
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)
    if border_color:
        borders = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            b = OxmlElement(f'w:{side}')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), '4')
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), border_color)
            borders.append(b)
        tblPr.append(borders)
    mar = OxmlElement('w:tblCellMar')
    for side, w in (('top', 57), ('left', 91), ('bottom', 57), ('right', 91)):
        m = OxmlElement(f'w:{side}')
        m.set(qn('w:w'), str(w))
        m.set(qn('w:type'), 'dxa')
        mar.append(m)
    tblPr.append(mar)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(col_cms):
                cell.width = Cm(col_cms[i])
    for i, w in enumerate(col_cms):
        if i < len(table.columns):
            table.columns[i].width = Cm(w)
    if header_repeat:
        trPr = table.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement('w:tblHeader')
        th.set(qn('w:val'), 'true')
        trPr.append(th)

def row_h(row, cm):
    row.height = Cm(cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

def fmt(p, size=10.5, bold=False, italic=False, color=TEXT, before=0, after=6,
        align=None, keep_next=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    pf.keep_with_next = keep_next
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = hexrgb(color)
        r.font.name = 'Calibri'
    return p

def cpara(cell, text='', **kw):
    if cell.paragraphs and not cell.paragraphs[0].runs and not cell.paragraphs[0].text:
        p = cell.paragraphs[0]
        if text:
            p.add_run(text)
    else:
        p = cell.add_paragraph(text)
    return fmt(p, **kw)

# ---------- documento ----------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--saida', default='Formulário de Levantamento de Processo - Base para POP - Grupo Gadens.docx')
    ap.add_argument('--logo', default=LOGO_PADRAO)
    args = ap.parse_args()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.5)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = hexrgb(TEXT)
    rf = normal.element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), 'Calibri')
    rf.set(qn('w:hAnsi'), 'Calibri')

    doc.core_properties.title = 'Formulário de Levantamento de Processo — Grupo Gadens'
    doc.core_properties.author = 'Grupo Gadens — Processos'

    # cabeçalho: logo + título
    ht = doc.add_table(rows=1, cols=2)
    tbl_setup(ht, [3.6, 13.4], border_color=None)
    lc, rc = ht.rows[0].cells
    lc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    rc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    lp = lc.paragraphs[0]
    if os.path.exists(args.logo):
        lp.add_run().add_picture(args.logo, width=Cm(3.0))
        fmt(lp, after=0)
    else:
        cpara(lc, 'GRUPO GADENS', size=13, bold=True, color=DARK, after=0)
    cpara(rc, 'FORMULÁRIO DE LEVANTAMENTO DE PROCESSO', size=15, bold=True,
          color=DARK, after=2)
    p2 = rc.add_paragraph('Conte como o seu trabalho é feito — a equipe de Processos transforma em POP')
    fmt(p2, size=9.5, italic=True, color=MIDGRAY, after=0)

    sep = doc.add_paragraph()
    para_border(sep, 'bottom', RED, sz=18)
    fmt(sep, size=2, after=10)

    # caixa de instruções
    it = doc.add_table(rows=1, cols=1)
    tbl_setup(it, [17.0])
    c = it.rows[0].cells[0]
    cell_bg(c, FILL_LIGHT)
    cpara(c, 'Antes de começar', size=11, bold=True, color=RED, after=4)
    cpara(c, 'Você conhece o seu trabalho melhor do que ninguém. Este formulário serve para '
             'você contar, com as suas palavras, como uma tarefa ou rotina é feita hoje. Com '
             'as respostas, a equipe de Processos escreve o POP (Procedimento Operacional '
             'Padrão) — o documento oficial com o passo a passo da atividade.', size=10, after=6)
    for t in ['Escreva como se estivesse explicando o trabalho a um colega novo no setor.',
              'Não se preocupe com termos técnicos nem com "escrever bonito" — vale escrever à mão ou digitar.',
              'Se uma pergunta não se aplicar, escreva "não se aplica". Se não souber, deixe em branco e siga.',
              'As linhas cinzas nas tabelas são apenas exemplos de preenchimento.',
              'Se quiser, anexe prints de tela, planilhas ou fotos que ajudem a entender.']:
        p = c.add_paragraph(t, style='List Bullet')
        fmt(p, size=10, after=2)

    fmt(doc.add_paragraph(), size=4, after=4)

    # identificação
    p = doc.add_paragraph('IDENTIFICAÇÃO')
    para_shading(p, DARK)
    para_left_border(p, RED)
    fmt(p, size=12, bold=True, color='FFFFFF', before=6, after=6, keep_next=True)
    idt = doc.add_table(rows=6, cols=2)
    tbl_setup(idt, [5.5, 11.5])
    for i, (lab, hint) in enumerate([
            ('Seu nome', ''), ('Seu cargo ou função', ''), ('Setor / área', ''),
            ('Empresa do grupo', ''), ('Data de preenchimento', ''),
            ('Nome da tarefa ou rotina descrita', 'dê o nome pelo qual ela é conhecida no dia a dia')]):
        row = idt.rows[i]
        row_h(row, 0.8)
        a, b = row.cells
        a.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_bg(a, FILL_LIGHT)
        cpara(a, lab, size=10, bold=True, color=DARK, after=0)
        if hint:
            cpara(b, hint, size=8.5, italic=True, color='A6A6A6', after=0)

    def section_bar(num, title, tag):
        p = doc.add_paragraph()
        para_shading(p, DARK)
        para_left_border(p, RED)
        pf = p.paragraph_format
        pf.tab_stops.add_tab_stop(Cm(16.7), WD_TAB_ALIGNMENT.RIGHT)
        r1 = p.add_run(f'{num}.  {title}')
        r1.font.size, r1.font.bold = Pt(12), True
        r1.font.color.rgb = hexrgb('FFFFFF')
        r1.font.name = 'Calibri'
        r2 = p.add_run(f'\t{tag}')
        r2.font.size = Pt(7.5)
        r2.font.color.rgb = hexrgb(TAG)
        r2.font.name = 'Calibri'
        pf.space_before, pf.space_after = Pt(14), Pt(6)
        pf.keep_with_next = True

    def question(text, hint=None, keep=True):
        fmt(doc.add_paragraph(text), size=10.8, bold=True, after=2 if hint else 5, keep_next=keep)
        if hint:
            fmt(doc.add_paragraph(hint), size=9, italic=True, color=MIDGRAY, after=5, keep_next=keep)

    def answer_box(height_cm):
        t = doc.add_table(rows=1, cols=1)
        tbl_setup(t, [17.0])
        row_h(t.rows[0], height_cm)
        fmt(t.rows[0].cells[0].paragraphs[0], after=0)

    def form_table(headers, col_cms, example, n_empty, empty_h=0.85, numbered=False):
        t = doc.add_table(rows=1 + (1 if example else 0) + n_empty, cols=len(headers))
        tbl_setup(t, col_cms, header_repeat=True)
        hr = t.rows[0]
        row_h(hr, 0.65)
        for j, h in enumerate(headers):
            cell = hr.cells[j]
            cell_bg(cell, DARK)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cpara(cell, h, size=9.8, bold=True, color='FFFFFF', after=0,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        off = 1
        if example:
            er = t.rows[1]
            row_h(er, 0.7)
            for j, txt in enumerate(example):
                cell = er.cells[j]
                cell_bg(cell, FILL_EX)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cpara(cell, txt, size=9, italic=True, color='909090', after=0)
            off = 2
        for i in range(n_empty):
            row = t.rows[off + i]
            row_h(row, empty_h)
            if numbered:
                cell = row.cells[0]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cpara(cell, str(i + 1), size=10, color='A6A6A6', after=0,
                      align=WD_ALIGN_PARAGRAPH.CENTER)

    def subhead(text):
        fmt(doc.add_paragraph(text), size=10.8, bold=True, color=RED, before=8, after=3, keep_next=True)

    section_bar(1, 'PARA QUE SERVE ESTE TRABALHO?', 'POP · Seção 1 — Objetivo')
    question('Explique em poucas linhas por que esse trabalho existe: o que ele entrega ou garante quando é bem feito?',
             'Ex.: "Garantir que os boletos sejam emitidos sem erro e enviados aos clientes até o dia 25 de cada mês."')
    answer_box(3.0)

    section_bar(2, 'ONDE ACONTECE E QUEM PARTICIPA', 'POP · Seção 2 — Aplicação')
    question('Em qual(is) empresa(s) do grupo esse trabalho é feito? Quais setores participam ou usam o resultado?',
             'Ex.: "Feito no setor Financeiro da Construtora; Vendas manda as informações e a Contabilidade usa o resultado."')
    answer_box(2.4)

    section_bar(3, 'REGRAS E DOCUMENTOS QUE VOCÊ SEGUE', 'POP · Seção 3 — Referências')
    question('Existe alguma lei, norma, manual, contrato ou orientação interna que precisa ser seguida nesse trabalho?',
             'Cite o que lembrar, mesmo sem saber o nome exato. Ex.: manual do sistema, política da empresa, norma técnica, exigência de cliente, legislação.')
    answer_box(2.4)

    section_bar(4, 'SIGLAS E TERMOS DO DIA A DIA', 'POP · Seção 4 — Definições')
    question('Quais siglas, apelidos ou termos técnicos são usados nesse trabalho? O que cada um significa?')
    form_table(['Sigla ou termo', 'O que significa'], [4.5, 12.5],
               ['NF', 'Nota Fiscal — documento emitido a cada venda'], 4)

    section_bar(5, 'O QUE VOCÊ USA PARA TRABALHAR', 'POP · Seção 5 — Recursos necessários')
    subhead('5.1 — Sistemas e ferramentas')
    question('Quais sistemas, programas ou ferramentas você usa nesse trabalho?')
    form_table(['Sistema ou ferramenta', 'Para que você usa'], [5.5, 11.5],
               ['Excel', 'Controlar os pedidos do mês'], 4)
    subhead('5.2 — O que precisa chegar até você')
    question('Quais informações, documentos ou arquivos você precisa receber para o trabalho começar ou andar?')
    form_table(['O que você recebe', 'De quem vem', 'Quando / com que frequência'], [7.0, 5.0, 5.0],
               ['Planilha de pedidos', 'Setor de Vendas', 'Todo dia, de manhã'], 4)
    subhead('5.3 — Pessoas essenciais')
    question('Sem quais funções ou pessoas esse trabalho não anda? (cargos/funções, não precisa de nomes)')
    answer_box(2.0)

    section_bar(6, 'O PASSO A PASSO', 'POP · Seção 6 — Principais passos do processo')
    question('O que dá a largada nesse trabalho?',
             'Ex.: uma data do mês, um e-mail que chega, um pedido do chefe, um cliente que liga.')
    answer_box(1.3)
    question('Agora descreva o trabalho do início ao fim, na ordem em que as coisas acontecem.',
             'Comece cada passo com um verbo (conferir, emitir, enviar, lançar...). Use quantas linhas precisar; se faltar espaço, continue no verso ou em outra folha.')
    form_table(['Nº', 'O que é feito', 'Quem faz', 'Em qual sistema / lugar', 'O que sai dessa etapa'],
               [1.0, 5.8, 3.0, 3.4, 3.8],
               ['Ex.', 'Conferir os pedidos recebidos no dia', 'Assistente de vendas', 'E-mail + planilha', 'Lista de pedidos conferida'],
               12, empty_h=0.8, numbered=True)
    question('Como você sabe que o trabalho terminou? O que marca o fim?',
             'Ex.: "Quando todos os boletos aparecem como enviados no sistema."')
    answer_box(1.3)

    section_bar(7, 'QUEM FAZ E QUEM DECIDE O QUÊ', 'POP · Seção 7 — Papéis, responsabilidades e autonomia')
    question('Para cada função envolvida: o que ela faz nesse trabalho e o que ela pode decidir sozinha? Quando precisa pedir aprovação, e para quem?')
    form_table(['Função / cargo (área)', 'O que faz nesse trabalho', 'Decide sozinho × precisa de aprovação'],
               [4.3, 6.2, 6.5],
               ['Analista financeiro (Financeiro)', 'Emite os boletos e confere os valores',
                'Corrige valores até R$ 100; acima disso, aprovação do coordenador'], 4, empty_h=1.0)

    section_bar(8, 'COMO SABER SE ESTÁ INDO BEM', 'POP · Seção 8 — Indicadores')
    question('Existe algum número, meta ou prazo que é acompanhado? Como vocês percebem se o trabalho vai bem ou mal?',
             'Ex.: quantidade feita por mês, percentual de erros, prazo máximo de entrega, reclamações recebidas. Se não houver nada medido hoje, diga isso também.')
    answer_box(2.4)

    section_bar(9, 'O QUE NÃO PODE DAR ERRADO', 'POP · Seção 9 — Pontos críticos do processo')
    question('Quais partes desse trabalho exigem mais atenção? Onde um erro causa mais prejuízo, retrabalho ou dor de cabeça?',
             'Ex.: "Errar o CNPJ na nota obriga a cancelar e refazer tudo"; "Se passar do dia 25, o cliente paga juros e reclama."')
    answer_box(2.8)

    section_bar(10, 'QUANDO ALGO SAI DO PREVISTO', 'POP · Seção 10 — Desvios / ações necessárias')
    question('O que costuma dar errado ou fugir do padrão nesse trabalho? O que é feito quando isso acontece, e quem resolve?')
    form_table(['O que pode dar errado', 'O que é feito quando acontece', 'Quem resolve'],
               [5.5, 7.0, 4.5],
               ['Sistema fora do ar', 'Registrar os pedidos em planilha e lançar depois', 'Assistente + coordenador'],
               4, empty_h=1.0)

    section_bar(11, 'O QUE FICA REGISTRADO OU GUARDADO', 'POP · Seção 11 — Registro do processo')
    question('O que esse trabalho gera no final (documentos, relatórios, cadastros, comprovantes)? Onde cada item fica guardado?')
    form_table(['O que é gerado', 'Onde fica guardado (pasta, sistema, arquivo físico)'],
               [7.0, 10.0],
               ['Boleto emitido', 'Sistema TOTVS e pasta do cliente no OneDrive'], 4)

    section_bar(12, 'ESPAÇO LIVRE', 'Informação de apoio para o elaborador do POP')
    question('Tem mais alguma coisa que quem for escrever o POP precisa saber?',
             'Ex.: exceções raras, épocas de pico, coisas que já deram errado no passado, ideias de melhoria, o que você mudaria se pudesse.')
    answer_box(3.2)

    ft = doc.add_table(rows=1, cols=1)
    tbl_setup(ft, [17.0])
    c = ft.rows[0].cells[0]
    cell_bg(c, FILL_LIGHT)
    cpara(c, 'Terminou?', size=10.5, bold=True, color=RED, after=2)
    cpara(c, 'Releia rapidinho para ver se ficou algo em branco que você saiba responder e entregue este '
             'formulário à equipe de Processos. O código do documento, o fluxo no Bizagi, as aprovações e o '
             'controle de emissão do POP serão preenchidos pela equipe de Processos/Qualidade — você não '
             'precisa se preocupar com isso.', size=9.5, after=0)

    # rodapé com paginação
    fp = sec.footer.paragraphs[0]
    fp.text = ''
    pf = fp.paragraph_format
    pf.tab_stops.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT)
    pPr = fp._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), '4')
    top.set(qn('w:space'), '3')
    top.set(qn('w:color'), 'D9D9D9')
    pBdr.append(top)
    pPr.append(pBdr)

    def foot_run(text):
        r = fp.add_run(text)
        r.font.size, r.font.color.rgb, r.font.name = Pt(8.5), hexrgb(MIDGRAY), 'Calibri'

    def foot_field(instr):
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), instr)
        fr = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '17'); rPr.append(sz)
        col = OxmlElement('w:color'); col.set(qn('w:val'), MIDGRAY); rPr.append(col)
        fr.append(rPr)
        t = OxmlElement('w:t'); t.text = '1'
        fr.append(t)
        fld.append(fr)
        fp._p.append(fld)

    foot_run('Formulário de Levantamento de Processo — Grupo Gadens\t')
    foot_run('Página ')
    foot_field(' PAGE ')
    foot_run(' de ')
    foot_field(' NUMPAGES ')

    doc.save(args.saida)
    print('OK:', args.saida)


if __name__ == '__main__':
    main()
