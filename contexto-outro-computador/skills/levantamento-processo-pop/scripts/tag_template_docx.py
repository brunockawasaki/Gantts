# -*- coding: utf-8 -*-
"""
Cria uma copia marcada (com tokens {{...}}) do template oficial de POP,
preservando a formatacao original. O arquivo oficial NUNCA e alterado.

Uso:
  python tag_template_docx.py --origem "<docx original>" --saida "<docx marcado>"
"""
import argparse
import docx
from docx.oxml.ns import qn


def _apply_default_format(run):
    """Aplica a formatacao padrao do documento (Arial 10pt preto) a um run novo.
    Necessario porque paragrafos/celulas vazios nao tem rPr proprio e cairiam no
    default do tema do Word (11pt, fonte do tema) em vez do padrao usado no resto
    do documento. Ver memoria: doc default e 11pt, nao 10pt -- sempre declarar sz/szCs."""
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.makeelement(qn('w:rFonts'), {
        qn('w:ascii'): 'Arial', qn('w:hAnsi'): 'Arial', qn('w:cs'): 'Arial',
    })
    rPr.append(rFonts)
    color = rPr.makeelement(qn('w:color'), {qn('w:val'): '000000', qn('w:themeColor'): 'text1'})
    rPr.append(color)
    sz = rPr.makeelement(qn('w:sz'), {qn('w:val'): '20'})
    rPr.append(sz)
    szCs = rPr.makeelement(qn('w:szCs'), {qn('w:val'): '20'})
    rPr.append(szCs)


def set_paragraph_token(paragraph, token):
    """Substitui o texto de um paragrafo por um unico token. Se o paragrafo ja
    tem runs (frase de instrucao original), reaproveita o primeiro run e sua
    formatacao. Se estava vazio (item de lista placeholder), cria um run novo
    com a formatacao padrao do documento (Arial 10pt)."""
    if paragraph.runs:
        paragraph.runs[0].text = token
        for r in paragraph.runs[1:]:
            r.text = ""
        return
    run = paragraph.add_run(token)
    _apply_default_format(run)


def set_cell_token(cell, token):
    """Preenche a primeira celula/paragrafo vazio de uma celula de tabela com o token."""
    p = cell.paragraphs[0]
    set_paragraph_token(p, token)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--origem", required=True)
    ap.add_argument("--saida", required=True)
    args = ap.parse_args()

    d = docx.Document(args.origem)
    paras = d.paragraphs
    tables = d.tables

    assert len(paras) == 54, f"Estrutura do template mudou (paragrafos={len(paras)}, esperado 54). Revise os indices antes de continuar."
    assert len(tables) == 5, f"Estrutura do template mudou (tabelas={len(tables)}, esperado 5). Revise os indices antes de continuar."

    # --- Paragrafos de texto corrido (valor unico) ---
    set_paragraph_token(paras[6], "{{OBJETIVO}}")
    set_paragraph_token(paras[8], "{{APLICACAO}}")
    set_paragraph_token(paras[10], "{{REFERENCIAS}}")

    # --- Itens de lista com marcador (repetiveis) ---
    set_paragraph_token(paras[17], "{{SISTEMA_ITEM}}")
    set_paragraph_token(paras[21], "{{ENTRADA_ITEM}}")
    set_paragraph_token(paras[25], "{{PAPEL_CHAVE_ITEM}}")
    set_paragraph_token(paras[30], "{{PASSO_MACRO_ITEM}}")
    set_paragraph_token(paras[39], "{{INDICADOR_ITEM}}")
    set_paragraph_token(paras[44], "{{PONTO_CRITICO_ITEM}}")

    # --- Tabela 0: Definicoes (Sigla | O que significa) ---
    row = tables[0].rows[1]
    set_cell_token(row.cells[0], "{{SIGLA}}")
    set_cell_token(row.cells[1], "{{SIGLA_DEF}}")

    # --- Tabela 1: Papeis e responsabilidades (Funcao | O que faz | O que decide) ---
    row = tables[1].rows[1]
    set_cell_token(row.cells[0], "{{FUNCAO}}")
    set_cell_token(row.cells[1], "{{FUNCAO_FAZ}}")
    set_cell_token(row.cells[2], "{{FUNCAO_DECIDE}}")

    # --- Tabela 2: Desvios / acoes (Problema | Acao | Responsavel) ---
    row = tables[2].rows[1]
    set_cell_token(row.cells[0], "{{DESVIO_PROBLEMA}}")
    set_cell_token(row.cells[1], "{{DESVIO_ACAO}}")
    set_cell_token(row.cells[2], "{{DESVIO_RESP}}")

    # --- Tabela 3: Registro do processo (Identificacao | Armazenamento) ---
    row = tables[3].rows[1]
    set_cell_token(row.cells[0], "{{REGISTRO_ITEM}}")
    set_cell_token(row.cells[1], "{{REGISTRO_LOCAL}}")

    # --- Tabela 4: Controle de emissao -> NAO tocar (equipe de Qualidade preenche) ---

    d.save(args.saida)
    print("OK ->", args.saida)


if __name__ == "__main__":
    main()
